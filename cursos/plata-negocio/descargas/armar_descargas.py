"""Arma las descargas de Llévatelo: plantillas Excel (U2, U3) y la hoja del cuaderno para imprimir (U3, HTML → PDF).
Uso: python3 armar_descargas.py && node ../guias-pdf/renderizar_pdf.mjs descargas"""
import pathlib
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

AQUI = pathlib.Path(__file__).parent
ROJO, NEGRO, GRIS, SUAVE = "C8102E", "141414", "8D8D94", "F6F6F7"
PESOS = '"$"#,##0;[Red]-"$"#,##0'
fino = Side(style="thin", color="D4D4D8")
borde = Border(left=fino, right=fino, top=fino, bottom=fino)


def cabecera(ws, titulo, sub, ancho):
    ws.sheet_view.showGridLines = False
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=ancho)
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=ancho)
    ws["A1"] = titulo
    ws["A1"].font = Font(bold=True, size=16, color="FFFFFF")
    ws["A2"] = sub
    ws["A2"].font = Font(size=10, color="FFFFFF")
    for r in (1, 2):
        for c in range(1, ancho + 1):
            ws.cell(r, c).fill = PatternFill("solid", fgColor=ROJO)
    ws.row_dimensions[1].height = 30
    ws.row_dimensions[2].height = 22


def encabezados(ws, fila, nombres):
    for c, n in enumerate(nombres, 1):
        x = ws.cell(fila, c, n)
        x.font = Font(bold=True, color="FFFFFF")
        x.fill = PatternFill("solid", fgColor=NEGRO)
        x.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        x.border = borde
    ws.row_dimensions[fila].height = 30


def entrada(c):
    c.fill = PatternFill("solid", fgColor="FFF7D6")
    c.border = borde


def calculo(c, negrita=False):
    c.fill = PatternFill("solid", fgColor=SUAVE)
    c.border = borde
    c.font = Font(bold=negrita)


# ---------- U2: Calcula tu sueldo ----------
wb = Workbook()
for nombre, ejemplo in (("Tu sueldo", False), ("Ejemplo Luz Dary", True)):
    ws = wb.active if not ejemplo else wb.create_sheet(nombre)
    ws.title = nombre
    cabecera(ws, "Calcula tu sueldo", "Unidad 2 · La plata del negocio no es la de la casa · Llena solo las celdas amarillas", 4)
    ws.column_dimensions["A"].width = 30
    for col in "BCD":
        ws.column_dimensions[col].width = 20
    encabezados(ws, 4, ["Mes", "Ventas del mes", "Pagos del negocio\n(mercancía, local, transporte)", "Ganancia del mes"])
    for i in range(3):
        f = 5 + i
        ws.cell(f, 1, f"Mes {i + 1}" if not ejemplo else ["Mes 1", "Mes 2", "Mes 3"][i])
        entrada(ws.cell(f, 1))
        for c in (2, 3):
            entrada(ws.cell(f, c))
            ws.cell(f, c).number_format = PESOS
        g = ws.cell(f, 4)
        if ejemplo:
            g.value = [2400000, 1800000, 2100000][i]
            entrada(g)
        else:
            g.value = f'=IF(AND(B{f}="",C{f}=""),"",B{f}-C{f})'
            calculo(g)
        g.number_format = PESOS
    filas = [("Promedio de 3 meses", "=IFERROR(AVERAGE(D5:D7),\"\")", True),
             ("Colchón (regla práctica del curso)", 0.2, False),
             ("Tu sueldo al mes", "=IFERROR(B9*(1-B10),\"\")", True),
             ("Por quincena", "=IFERROR(B11/2,\"\")", True)]
    for k, (et, v, neg) in enumerate(filas):
        f = 9 + k
        ws.cell(f, 1, et).font = Font(bold=True)
        c = ws.cell(f, 2, v)
        if k == 1:
            entrada(c)
            c.number_format = "0%"
        else:
            calculo(c, neg)
            c.number_format = PESOS
    ws["B11"].fill = PatternFill("solid", fgColor=ROJO)
    ws["B11"].font = Font(bold=True, color="FFFFFF", size=12)
    notas = ["Cómo se usa:",
             "1. Escribe las ventas y lo que pagaste para el negocio en cada uno de los últimos 3 meses.",
             "2. La hoja calcula la ganancia de cada mes, el promedio y tu sueldo con el colchón.",
             "3. El colchón del 20 % es una regla práctica del curso, no una norma: tu contador puede ajustarlo.",
             "4. Escoge tus fechas fijas de pago (por ejemplo, el 15 y el 30) y págate por transferencia.",
             "Si el sueldo no alcanza para la casa, es una señal: vender más, revisar precios o bajar costos."]
    if ejemplo:
        notas[1] = "En este ejemplo solo se conoce la ganancia de cada mes de Luz Dary (Unidad 2)."
    for k, t in enumerate(notas):
        c = ws.cell(14 + k, 1, t)
        c.font = Font(bold=(k == 0), color="45454C" if k else NEGRO, size=10)
        ws.merge_cells(start_row=14 + k, start_column=1, end_row=14 + k, end_column=4)
wb.save(AQUI / "u2-calcula-tu-sueldo.xlsx")

# ---------- U3: Registro diario ----------
wb = Workbook()
for nombre, ejemplo in (("Registro", False), ("Ejemplo", True)):
    ws = wb.active if not ejemplo else wb.create_sheet(nombre)
    ws.title = nombre
    cabecera(ws, "Registro diario del negocio", "Unidad 3 · 5 minutos al cerrar · Saldo de hoy = saldo de ayer + entró − salió", 5)
    for col, w in zip("ABCDE", (14, 36, 16, 16, 18)):
        ws.column_dimensions[col].width = w
    encabezados(ws, 4, ["Fecha", "Qué fue", "Entró", "Salió", "Saldo"])
    ws.cell(5, 2, "Saldo con el que empiezas")
    ws.cell(5, 2).font = Font(italic=True)
    for c in (1, 2, 3, 4):
        entrada(ws.cell(5, c))
    s = ws.cell(5, 5)
    entrada(s)
    s.number_format = PESOS
    datos = []
    if ejemplo:
        s.value = 400000
        ws.cell(5, 1, "Lunes")
        datos = [("Lunes", "Ventas del día", 350000, None), ("Lunes", "Pago al proveedor", None, 120000)]
    total = 40 if not ejemplo else 2
    for i in range(total):
        f = 6 + i
        for c in (1, 2, 3, 4):
            entrada(ws.cell(f, c))
        ws.cell(f, 1).number_format = "dd/mm/yyyy"
        for c in (3, 4):
            ws.cell(f, c).number_format = PESOS
        if i < len(datos):
            for c, v in zip((1, 2, 3, 4), datos[i]):
                if v is not None:
                    ws.cell(f, c, v)
        e = ws.cell(f, 5, f'=IF(AND(C{f}="",D{f}=""),"",LOOKUP(2,1/(E$5:E{f - 1}<>""),E$5:E{f - 1})+N(C{f})-N(D{f}))')
        calculo(e, True)
        e.number_format = PESOS
    if ejemplo:
        f = 6 + total + 1
        ws.cell(f, 2, "Saldo al cerrar (Luz Dary): $630.000").font = Font(bold=True, color=ROJO)
    ws.freeze_panes = "A5"
    fin = 6 + total + (3 if ejemplo else 1)
    for k, t in enumerate(["Si el saldo no cuadra con la cuenta y el cajón, algo se fue sin anotar.",
                           "No olvides anotar tu sueldo (sale de la cuenta del negocio) ni las ventas en efectivo."]):
        ws.cell(fin + k, 1, t).font = Font(size=10, color="45454C")
        ws.merge_cells(start_row=fin + k, start_column=1, end_row=fin + k, end_column=5)
wb.save(AQUI / "u3-registro-diario.xlsx")

# ---------- U3: hoja del cuaderno para imprimir ----------
filas = "".join("<tr><td></td><td></td><td></td><td></td><td></td></tr>" for _ in range(22))
html = f"""<!doctype html><html lang="es"><head><meta charset="utf-8"><title>Registro diario</title><style>
@page{{size:A4;margin:12mm}}*{{box-sizing:border-box;margin:0}}
body{{font-family:Inter,'DejaVu Sans',sans-serif;color:#1a1a1a;-webkit-print-color-adjust:exact;print-color-adjust:exact}}
.cab{{display:flex;justify-content:space-between;align-items:flex-end;border-bottom:1.2mm solid #c8102e;padding-bottom:3mm}}
h1{{font-size:20pt;font-weight:800}}.cab small{{font-size:8pt;color:#6b6b72;text-transform:uppercase;letter-spacing:.08em;font-weight:700}}
.f{{margin:4mm 0;background:#141414;color:#fff;border-radius:2mm;padding:3mm 4mm;font-size:11pt}}.f b{{color:#ff8a9c}}
.datos{{display:flex;gap:6mm;font-size:10pt;margin-bottom:3mm}}.datos span{{flex:1;border-bottom:.3mm solid #8d8d94;padding-bottom:1mm}}
table{{width:100%;border-collapse:collapse;font-size:10pt}}th{{background:#141414;color:#fff;padding:2.2mm;font-weight:700}}
td{{border:.3mm solid #b8b8be;height:8.6mm}}th:nth-child(1){{width:16%}}th:nth-child(2){{width:36%}}
tr.ini td{{background:#f6f6f7}}.pie{{margin-top:3mm;font-size:8pt;color:#6b6b72;display:flex;justify-content:space-between}}
</style></head><body>
<div class="cab"><h1>Registro diario del negocio</h1><small>Unidad 3 · La plata del negocio no es la de la casa</small></div>
<div class="f"><b>Saldo de hoy</b> = saldo de ayer + lo que entró − lo que salió</div>
<div class="datos"><span>Negocio:</span><span>Mes:</span><span>Cuenta del negocio:</span></div>
<table><tr><th>Fecha</th><th>Qué fue</th><th>Entró</th><th>Salió</th><th>Saldo</th></tr>
<tr class="ini"><td></td><td><i>Saldo con el que empiezas</i></td><td></td><td></td><td></td></tr>{filas}</table>
<div class="pie"><span>Si el saldo no cuadra con la cuenta y el cajón, algo se fue sin anotar. Anota también tu sueldo.</span><span>Aula Virtual Empresarial Soacha</span></div>
</body></html>"""
(AQUI / "u3-hoja-registro.html").write_text(html, encoding="utf-8")
print("ok")
